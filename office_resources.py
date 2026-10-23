"""Explicit owner-selected document excerpts; no macros or embedded code."""
from pathlib import Path
import zipfile

def extract(path):
 path=Path(path)
 if path.stat().st_size>8*1048576:raise ValueError('Choose documents smaller than 8 MiB.')
 suffix=path.suffix.lower()
 if suffix in ('.docx','.xlsx'):
  with zipfile.ZipFile(path) as archive:
   if sum(item.file_size for item in archive.infolist())>40*1048576:raise ValueError('Expanded document exceeds 40 MiB.')
 if suffix in ('.txt','.csv','.md','.json'):return path.read_text(encoding='utf-8-sig',errors='replace')[:30000]
 if suffix=='.docx':
  from docx import Document
  doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
  for table in doc.tables:
   for row in table.rows:text+='\n'+' | '.join(cell.text for cell in row.cells)
  return text[:30000]
 if suffix=='.xlsx':
  from openpyxl import load_workbook
  book=load_workbook(path,read_only=True,data_only=True)
  try:
   text=[]
   for sheet in list(book)[:10]:
    text.append('Sheet: '+sheet.title)
    for row in sheet.iter_rows(max_row=500,max_col=30,values_only=True):text.append(' | '.join(str(value) if value is not None else '' for value in row))
   return '\n'.join(text)[:30000]
  finally:book.close()
 if suffix=='.pdf':
  from pypdf import PdfReader
  reader=PdfReader(path)
  if reader.is_encrypted:raise ValueError('Choose an unlocked PDF.')
  text='\n'.join(page.extract_text() or '' for page in list(reader.pages)[:30])
  if not text.strip():raise ValueError('This PDF needs OCR; supply readable text or a text PDF.')
  return text[:30000]
 raise ValueError('Use TXT, MD, CSV, JSON, DOCX, XLSX or a text PDF.')
