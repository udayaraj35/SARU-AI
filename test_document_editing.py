import tempfile,unittest,zipfile
from pathlib import Path
from docx import Document
from docx.shared import Pt
import pymupdf as fitz
from document_editing import replace_text,fingerprint

class EditTests(unittest.TestCase):
 def test_docx_run_properties_and_source_remain(self):
  with tempfile.TemporaryDirectory() as temp:
   source=Path(temp)/'original.docx';target=Path(temp)/'copy.docx';doc=Document();run=doc.add_paragraph().add_run('Hello Alice');run.font.name='Arial';run.font.size=Pt(14);run.bold=True;doc.save(source);before=source.read_bytes()
   replace_text(source,target,'Alice','Maya',expected=fingerprint(source));out=Document(target);r=out.paragraphs[0].runs[0]
   self.assertEqual(r.text,'Hello Maya');self.assertEqual(r.font.name,'Arial');self.assertEqual(r.font.size.pt,14);self.assertTrue(r.bold);self.assertEqual(source.read_bytes(),before)
   with zipfile.ZipFile(source) as a,zipfile.ZipFile(target) as b:
    for name in a.namelist():
     if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name))
 def test_docx_ambiguous_or_split_run_is_not_rebuilt(self):
  with tempfile.TemporaryDirectory() as temp:
   source=Path(temp)/'source.docx';target=Path(temp)/'out.docx';doc=Document();p=doc.add_paragraph();p.add_run('Ali');p.add_run('ce');doc.save(source)
   with self.assertRaisesRegex(ValueError,'single formatted'):replace_text(source,target,'Alice','Maya')
   self.assertFalse(target.exists())
 def test_pdf_font_size_and_other_content_preserved(self):
  with tempfile.TemporaryDirectory() as temp:
   source=Path(temp)/'source.pdf';target=Path(temp)/'out.pdf';doc=fitz.open();page=doc.new_page();page.insert_text((72,90),'Hello Alice',fontname='hebo',fontsize=16);page.insert_text((72,180),'Unchanged text',fontsize=10);doc.save(source);doc.close();before=source.read_bytes()
   replace_text(source,target,'Hello Alice','Hello Ana',expected=fingerprint(source))
   with fitz.open(target) as out:
    spans=[s for b in out[0].get_text('dict')['blocks'] for line in b.get('lines',[]) for s in line['spans']];edited=next(s for s in spans if s['text']=='Hello Ana')
    self.assertEqual(edited['size'],16);self.assertEqual(edited['font'],'Helvetica-Bold');self.assertIn('Unchanged text',out[0].get_text())
   self.assertEqual(source.read_bytes(),before)
 def test_pdf_overflow_and_source_change_refused_without_copy(self):
  with tempfile.TemporaryDirectory() as temp:
   source=Path(temp)/'source.pdf';target=Path(temp)/'out.pdf';doc=fitz.open();doc.new_page().insert_text((72,90),'Short',fontsize=12);doc.save(source);doc.close()
   with self.assertRaisesRegex(ValueError,'wider'):replace_text(source,target,'Short','Much much longer replacement')
   self.assertFalse(target.exists())
   with self.assertRaisesRegex(ValueError,'Source changed'):replace_text(source,target,'Short','Tiny',expected='incorrect')
   with self.assertRaisesRegex(ValueError,'original overwrite'):replace_text(source,source,'Short','Tiny')
if __name__=='__main__':unittest.main()
